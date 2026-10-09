# -*- coding: utf-8 -*-
"""
Exporteert alle teksten van de Calm Investments website naar een Word-document
(om in te werken) en een PDF (om te lezen of te printen).

Gebruik:  python scripts/export-teksten.py
De teksten worden live uit content/*.json gelezen, dus na elke wijziging
opnieuw draaien om een actueel document te krijgen.
"""
import json
from pathlib import Path
from datetime import date
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parent.parent
CONTENT = ROOT / "content"
OUT = ROOT / "exports"
OUT.mkdir(exist_ok=True)

nl = json.loads((CONTENT / "nl.json").read_text(encoding="utf-8"))
en = json.loads((CONTENT / "en.json").read_text(encoding="utf-8"))
pf = json.loads((CONTENT / "portfolio.json").read_text(encoding="utf-8"))

GREEN = (0x14 / 255, 0x50 / 255, 0x29 / 255)
GREEN_HEX = "145029"
MONTHS = ["januari", "februari", "maart", "april", "mei", "juni", "juli",
          "augustus", "september", "oktober", "november", "december"]
TODAY = date.today()
DATESTR = "%d %s %d" % (TODAY.day, MONTHS[TODAY.month - 1], TODAY.year)

sections = []


def sec(title, where, rows):
    sections.append({
        "num": len(sections) + 1,
        "title": title,
        "where": where,
        "rows": [r for r in rows if r and (r[1] or r[2])],
    })


def pair(label, path):
    """Haal dezelfde sleutel op uit het NL- en het EN-bestand."""
    def dig(src):
        cur = src
        for part in path.split("."):
            if part.endswith("]"):
                key, idx = part[:-1].split("[")
                cur = cur[key][int(idx)] if key else cur[int(idx)]
            else:
                cur = cur[part]
        return cur
    try:
        return (label, dig(nl), dig(en))
    except (KeyError, IndexError):
        return None


# ---------------------------------------------------------------- secties
sec("Menu", "De balk bovenaan elke pagina.", [
    pair("Menu-item 1", "nav.about"),
    pair("Menu-item 2", "nav.approach"),
    pair("Menu-item 3", "nav.portfolio"),
    pair("Menu-item 4", "nav.contact"),
])

sec("Openingsblok", "Het eerste wat een bezoeker ziet, bovenaan de homepage.", [
    pair("Klein kopje boven de titel", "hero.label"),
    pair("Grote titel", "hero.title"),
    pair("Introductietekst", "hero.subtitle"),
    pair("Groene knop", "hero.ctaPrimary"),
    pair("Link ernaast", "hero.ctaSecondary"),
])

sec("Over ons", "Het groene blok onder het openingsblok.", [
    pair("Titel", "about.title"),
    pair("Alinea 1", "about.paragraphs[0]"),
    pair("Alinea 2", "about.paragraphs[1]"),
    pair("Kerncijfer 1: getal", "about.stats[0].value"),
    pair("Kerncijfer 1: tekst eronder", "about.stats[0].label"),
    pair("Kerncijfer 2: getal", "about.stats[1].value"),
    pair("Kerncijfer 2: tekst eronder", "about.stats[1].label"),
])

approach_rows = [
    pair("Klein kopje", "approach.label"),
    pair("Titel", "approach.title"),
    pair("Introductietekst", "approach.intro"),
]
for i in range(len(nl["approach"]["pillars"])):
    approach_rows.append(pair("Pijler %d: kop" % (i + 1), "approach.pillars[%d].title" % i))
    approach_rows.append(pair("Pijler %d: tekst" % (i + 1), "approach.pillars[%d].text" % i))
sec("Aanpak", "De vier pijlers, elk met een icoon erboven.", approach_rows)

sec("Deelnemingen op de homepage", "Het groene blok met de drie bedrijven.", [
    pair("Titel", "portfolio.title"),
    pair("Tekst eronder", "portfolio.subtitle"),
    pair("Link op elke kaart", "portfolio.viewItem"),
    pair("Knop onderaan", "portfolio.viewAll"),
])

company_rows = []
for item in pf["items"]:
    name = item["name"]
    company_rows.append(("%s: naam" % name, name, name))
    company_rows.append(("%s: omschrijving" % name,
                         item["summary"]["nl"], item["summary"]["en"]))
    company_rows.append(("%s: labels" % name,
                         ", ".join(item["tags"]["nl"]), ", ".join(item["tags"]["en"])))
sec("De bedrijven zelf",
    "De kaarten met de deelnemingen. Bedrijven toevoegen of weghalen kan ook, "
    "geef dat hieronder aan.",
    company_rows)

focus_rows = [
    pair("Klein kopje", "focus.label"),
    pair("Titel", "focus.title"),
    pair("Introductietekst", "focus.intro"),
]
for i in range(len(nl["focus"]["areas"])):
    focus_rows.append(pair("Focusgebied %d" % (i + 1), "focus.areas[%d]" % i))
focus_rows.append(pair("Kop boven de criteria", "focus.criteriaTitle"))
for i in range(len(nl["focus"]["criteria"])):
    focus_rows.append(pair("Criterium %d" % (i + 1), "focus.criteria[%d]" % i))
sec("Focusgebieden", "De genummerde lijst met sectoren en de investeringscriteria.", focus_rows)

team_rows = [pair("Titel van het blok", "team.title")]
for i in range(len(nl["team"]["members"])):
    naam = nl["team"]["members"][i]["name"]
    team_rows.append(pair("%s: naam" % naam, "team.members[%d].name" % i))
    team_rows.append(pair("%s: functie" % naam, "team.members[%d].role" % i))
    team_rows.append(pair("%s: biografie" % naam, "team.members[%d].bio" % i))
    team_rows.append(pair("%s: persoonlijke regel" % naam, "team.members[%d].personality" % i))
sec("Team", "Het groene blok met de portretten van Olav en Jonathan.", team_rows)

sec("Contact", "Het laatste blok van de homepage, met het contactformulier.", [
    pair("Titel", "contact.title"),
    pair("Tekst eronder", "contact.text"),
    pair("Naam van het pand", "contact.addressLabel"),
    pair("Adres", "contact.address"),
    pair("Telefoonnummer", "contact.phone"),
    pair("E-mailadres", "contact.email"),
    pair("Formulier: label naam", "contact.form.name"),
    pair("Formulier: voorbeeldtekst naam", "contact.form.namePlaceholder"),
    pair("Formulier: label e-mail", "contact.form.email"),
    pair("Formulier: voorbeeldtekst e-mail", "contact.form.emailPlaceholder"),
    pair("Formulier: label bericht", "contact.form.message"),
    pair("Formulier: voorbeeldtekst bericht", "contact.form.messagePlaceholder"),
    pair("Formulier: knop", "contact.form.submit"),
    pair("Bevestiging na verzenden", "contact.form.success"),
])

sec("Footer", "De groene balk onderaan elke pagina.", [
    pair("Tekst onder het logo", "footer.description"),
    pair("Kop boven de menulinks", "footer.navigateLabel"),
    pair("Kop boven de juridische links", "footer.legalLabel"),
    pair("Copyright", "footer.copyright"),
    pair("Slogan rechtsonder", "footer.tagline"),
])

sec("Pagina Deelnemingen", "De aparte pagina achter de knop onderaan het groene blok.", [
    pair("Klein kopje", "portfolio.pageLabel"),
    pair("Titel", "portfolio.pageTitle"),
    pair("Introductietekst", "portfolio.pageIntro"),
    pair("Link terug", "portfolio.backHome"),
])

sec("Cookiemelding",
    "De balk onderaan bij het eerste bezoek. Verschijnt pas zodra Google "
    "Analytics wordt aangezet.", [
    pair("Tekst", "cookieBanner.text"),
    pair("Knop accepteren", "cookieBanner.accept"),
    pair("Knop weigeren", "cookieBanner.reject"),
    pair("Link naar privacyverklaring", "cookieBanner.more"),
    pair("Link in de footer", "cookieBanner.settings"),
])

sec("Meldingen in het contactformulier",
    "Deze ziet een bezoeker alleen als er iets misgaat bij het invullen.", [
    pair("Naam ontbreekt", "contact.form.errors.nameRequired"),
    pair("E-mail ontbreekt", "contact.form.errors.emailRequired"),
    pair("E-mail klopt niet", "contact.form.errors.emailInvalid"),
    pair("Bericht ontbreekt", "contact.form.errors.messageRequired"),
    pair("Invoer te lang", "contact.form.errors.tooLong"),
    pair("Beveiligingscontrole", "contact.form.errors.captchaRequired"),
    pair("Verzenden mislukt", "contact.form.errors.generic"),
    pair("Te veel pogingen", "contact.form.errors.rateLimited"),
])

sec("Foutpagina s",
    "Wat een bezoeker ziet bij een verkeerde link of bij een storing.", [
    pair("Pagina niet gevonden: titel", "errors.notFoundTitle"),
    pair("Pagina niet gevonden: tekst", "errors.notFoundText"),
    pair("Storing: titel", "errors.errorTitle"),
    pair("Storing: tekst", "errors.errorText"),
    pair("Knop opnieuw proberen", "errors.retry"),
    pair("Knop terug naar home", "errors.backHome"),
])

sec("Google en social media",
    "Niet zichtbaar op de site zelf. Dit is de tekst die Google in de "
    "zoekresultaten toont en die verschijnt zodra iemand de link deelt.", [
    pair("Titel in Google", "meta.title"),
    pair("Omschrijving in Google", "meta.description"),
])

# ------------------------------------------------- juridische bijlage
legal_sections = []
for key, naam in (("privacy", "Privacyverklaring"), ("terms", "Algemene voorwaarden")):
    rows = [("Titel", nl["legalPages"][key]["title"], en["legalPages"][key]["title"]),
            ("Datum", nl["legalPages"][key]["updated"], en["legalPages"][key]["updated"])]
    for i, blok in enumerate(nl["legalPages"][key]["sections"]):
        blok_en = en["legalPages"][key]["sections"][i]
        rows.append(("Kop", blok["title"], blok_en["title"]))
        for j, alinea in enumerate(blok["paragraphs"]):
            alinea_en = blok_en["paragraphs"][j] if j < len(blok_en["paragraphs"]) else ""
            rows.append(("Alinea %d" % (j + 1), alinea, alinea_en))
    legal_sections.append({"title": naam, "rows": rows})

INTRO = [
    "Dit document bevat alle teksten die op de nieuwe website van Calm Investments staan, "
    "in dezelfde volgorde als waarin ze op de site verschijnen. Zo is per tekst te zien waar "
    "die terechtkomt.",
    "De site is tweetalig. In de middelste kolom staat de Nederlandse tekst, in de rechter "
    "kolom de Engelse. De Engelse teksten zijn gebaseerd op de huidige website; de Nederlandse "
    "zijn daarvan een vertaling. Beide mogen uiteraard aangepast worden.",
    "Aanpassen kan rechtstreeks in dit document: typ de gewenste tekst over de bestaande heen, "
    "of zet er een opmerking bij. Uitprinten en met pen corrigeren werkt ook prima, net als "
    "alleen de punten die opvallen in een mail zetten. Wat het makkelijkst is.",
    "De teksten hoeven niet in een keer af te zijn. Ook losse opmerkingen zoals dit klopt niet "
    "of hier mist iets over onze werkwijze zijn goed bruikbaar.",
]

INTRO_NOTE = ("De juridische teksten (privacyverklaring en algemene voorwaarden) staan "
              "achteraan als bijlage. Daarbij gaat het er vooral om of de feitelijke gegevens "
              "kloppen. Voor de juridische inhoud is het verstandig een jurist mee te laten lezen.")

HEADERS = ("Waar op de site", "Nederlands", "English")
TITEL = "Teksten website Calm Investments"


# ------------------------------------------------------------------ PDF
def build_pdf(path):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                    Table, TableStyle, PageBreak)

    def safe(text):
        repl = {"‘": "'", "’": "'", "“": '"', "”": '"',
                "–": "-", "—": "-", "€": "EUR ", "•": "-",
                " ": " "}
        text = str(text)
        for a, b in repl.items():
            text = text.replace(a, b)
        return escape(text)

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontName="Helvetica",
                          fontSize=8.5, leading=11.5)
    label = ParagraphStyle("label", parent=body,
                           textColor=colors.Color(0.30, 0.33, 0.29))
    head = ParagraphStyle("head", parent=body, fontName="Helvetica-Bold",
                          fontSize=8, textColor=colors.white)
    h1 = ParagraphStyle("h1", parent=styles["Normal"], fontName="Helvetica-Bold",
                        fontSize=15, leading=19, spaceAfter=2,
                        textColor=colors.Color(*GREEN))
    where = ParagraphStyle("where", parent=body, fontSize=9, leading=12.5,
                           textColor=colors.Color(0.35, 0.38, 0.35), spaceAfter=7)
    intro = ParagraphStyle("intro", parent=body, fontSize=10, leading=15, spaceAfter=9)

    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4),
                            leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=16 * mm,
                            title=TITEL, author="Luuk van Dijk")
    width = doc.width
    cols = [width * 0.19, width * 0.405, width * 0.405]
    story = []

    title_style = ParagraphStyle("title", parent=styles["Normal"],
                                 fontName="Helvetica-Bold", fontSize=24, leading=28,
                                 textColor=colors.Color(*GREEN))
    story.append(Paragraph(safe(TITEL), title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(safe("Versie %s" % DATESTR), where))
    story.append(Spacer(1, 10))
    for para in INTRO:
        story.append(Paragraph(safe(para), intro))
    note = ParagraphStyle("note", parent=intro, textColor=colors.Color(*GREEN))
    story.append(Paragraph(safe(INTRO_NOTE), note))

    sec_title = ParagraphStyle("sectitle", parent=body, fontName="Helvetica-Bold",
                               fontSize=12.5, leading=15, textColor=colors.white)
    sec_where = ParagraphStyle("secwhere", parent=body, fontSize=8.5, leading=11,
                               textColor=colors.Color(0.78, 0.84, 0.77), spaceBefore=2)

    def table_for(rows, title, where_text):
        banner = [Paragraph(safe(title), sec_title)]
        if where_text:
            banner.append(Paragraph(safe(where_text), sec_where))
        data = [[banner, "", ""],
                [Paragraph(safe(h), head) for h in HEADERS]]
        for lab, a, b in rows:
            data.append([Paragraph(safe(lab), label),
                         Paragraph(safe(a), body),
                         Paragraph(safe(b), body)])
        t = Table(data, colWidths=cols, repeatRows=2)
        t.setStyle(TableStyle([
            ("SPAN", (0, 0), (-1, 0)),
            ("BACKGROUND", (0, 0), (-1, 1), colors.Color(*GREEN)),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.Color(0.45, 0.6, 0.5)),
            ("TOPPADDING", (0, 0), (-1, 0), 8),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 1), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("GRID", (0, 1), (-1, -1), 0.4, colors.Color(0.82, 0.81, 0.78)),
            ("BOX", (0, 0), (-1, 0), 0.4, colors.Color(*GREEN)),
            ("ROWBACKGROUNDS", (0, 2), (-1, -1),
             [colors.white, colors.Color(0.97, 0.97, 0.95)]),
        ]))
        return t

    for s in sections:
        story.append(Spacer(1, 18))
        story.append(table_for(s["rows"], "%d. %s" % (s["num"], s["title"]), s["where"]))

    story.append(PageBreak())
    story.append(Paragraph(safe("Bijlage: juridische teksten"), h1))
    story.append(Paragraph(safe(INTRO_NOTE), where))
    for ls in legal_sections:
        story.append(Spacer(1, 14))
        story.append(table_for(ls["rows"], ls["title"], ""))

    def footer(canvas, docu):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColorRGB(0.45, 0.47, 0.45)
        canvas.drawString(15 * mm, 9 * mm,
                          "Teksten website Calm Investments  |  %s" % DATESTR)
        canvas.drawRightString(landscape(A4)[0] - 15 * mm, 9 * mm, "Pagina %d" % docu.page)
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# ----------------------------------------------------------------- DOCX
def build_docx(path):
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    d = Document()
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Mm(297), Mm(210)
    s.left_margin = s.right_margin = Mm(15)
    s.top_margin = Mm(15)
    s.bottom_margin = Mm(15)

    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    def shade(cell, hexcolor):
        el = OxmlElement("w:shd")
        el.set(qn("w:val"), "clear")
        el.set(qn("w:fill"), hexcolor)
        cell._tc.get_or_add_tcPr().append(el)

    t = d.add_paragraph()
    run = t.add_run(TITEL)
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x14, 0x50, 0x29)
    sub = d.add_paragraph()
    r = sub.add_run("Versie %s" % DATESTR)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x5A, 0x60, 0x5A)

    for para in INTRO:
        p = d.add_paragraph(para)
        p.paragraph_format.space_after = Pt(8)
    note = d.add_paragraph()
    rn = note.add_run(INTRO_NOTE)
    rn.font.color.rgb = RGBColor(0x14, 0x50, 0x29)

    widths = (Mm(50), Mm(108), Mm(108))

    def repeat_header(row):
        tr_pr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        tr_pr.append(el)

    def add_table(rows):
        table = d.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.autofit = False
        repeat_header(table.rows[0])
        hdr = table.rows[0].cells
        for i, h in enumerate(HEADERS):
            hdr[i].text = ""
            hrun = hdr[i].paragraphs[0].add_run(h)
            hrun.font.bold = True
            hrun.font.size = Pt(9)
            hrun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade(hdr[i], GREEN_HEX)
        for lab, a, b in rows:
            cells = table.add_row().cells
            for i, val in enumerate((str(lab), str(a), str(b))):
                cells[i].text = ""
                crun = cells[i].paragraphs[0].add_run(val)
                crun.font.size = Pt(9)
                if i == 0:
                    crun.font.color.rgb = RGBColor(0x4C, 0x54, 0x4B)
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = widths[i]
        return table

    for sct in sections:
        h = d.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        hrun = h.add_run("%d. %s" % (sct["num"], sct["title"]))
        hrun.font.size = Pt(15)
        hrun.font.bold = True
        hrun.font.color.rgb = RGBColor(0x14, 0x50, 0x29)
        w = d.add_paragraph()
        wrun = w.add_run(sct["where"])
        wrun.font.size = Pt(9)
        wrun.font.color.rgb = RGBColor(0x5A, 0x60, 0x5A)
        add_table(sct["rows"])

    d.add_page_break()
    h = d.add_paragraph()
    hrun = h.add_run("Bijlage: juridische teksten")
    hrun.font.size = Pt(15)
    hrun.font.bold = True
    hrun.font.color.rgb = RGBColor(0x14, 0x50, 0x29)
    wn = d.add_paragraph()
    wnrun = wn.add_run(INTRO_NOTE)
    wnrun.font.size = Pt(9)
    wnrun.font.color.rgb = RGBColor(0x5A, 0x60, 0x5A)
    for ls in legal_sections:
        lh = d.add_paragraph()
        lh.paragraph_format.space_before = Pt(16)
        lrun = lh.add_run(ls["title"])
        lrun.font.size = Pt(13)
        lrun.font.bold = True
        lrun.font.color.rgb = RGBColor(0x14, 0x50, 0x29)
        add_table(ls["rows"])

    d.save(str(path))


if __name__ == "__main__":
    pdf_path = OUT / "Calm-Investments-teksten.pdf"
    docx_path = OUT / "Calm-Investments-teksten.docx"
    build_pdf(pdf_path)
    build_docx(docx_path)
    total = sum(len(s["rows"]) for s in sections) + sum(len(x["rows"]) for x in legal_sections)
    print("secties : %d" % len(sections))
    print("teksten : %d" % total)
    print("pdf     : %s" % pdf_path)
    print("docx    : %s" % docx_path)
